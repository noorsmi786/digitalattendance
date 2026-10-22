import io
import csv
from datetime import datetime
from zoneinfo import ZoneInfo

PKT = ZoneInfo("Asia/Karachi")
from typing import List, Dict, Any, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.database import get_db_connection
from app.qr_service import generate_qr_payload, compute_student_code_and_roll

def generate_student_excel_template() -> io.BytesIO:
    """Creates a downloadable template for bulk importing students"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students Import"

    headers = [
        "Full Name*", "Father Name", "Class*", "Roll No*", "Parent Phone", "Gender", "Address"
    ]
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Add sample rows
    samples = [
        ["Muhammad Ali", "Tariq Khan", "10th", "01", "+923001234567", "Male", "Loi Shalman"],
        ["Fatima Bibi", "Abdul Rehman", "9th", "02", "+923019876543", "Female", "Loi Shalman"],
        ["Hamza Khan", "Zahid Khan", "KG", "05", "+923023456789", "Male", "Loi Shalman"]
    ]
    for row in samples:
        ws.append(row)

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def import_students_from_excel(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Reads Excel (.xlsx) or CSV file and imports students into SQLite database"""
    students_to_insert = []
    errors = []

    if filename.lower().endswith(".csv"):
        text = file_bytes.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            return {"success": False, "message": "File is empty"}
        headers = [h.strip().lower() for h in rows[0]]
        data_rows = rows[1:]
    else:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            return {"success": False, "message": "Excel sheet is empty"}
        headers = [str(h or "").strip().lower() for h in all_rows[0]]
        data_rows = all_rows[1:]

    # Map column headers flexibly
    def find_col(aliases: List[str]) -> int:
        for idx, h in enumerate(headers):
            for a in aliases:
                if a in h:
                    return idx
        return -1

    name_idx = find_col(["full name", "name", "student name"])
    father_idx = find_col(["father", "father name", "guardian"])
    class_idx = find_col(["class", "grade"])
    roll_idx = find_col(["roll", "roll no", "roll number"])
    phone_idx = find_col(["phone", "parent phone", "contact", "mobile", "whatsapp"])
    gender_idx = find_col(["gender", "sex"])
    address_idx = find_col(["address", "location", "city"])

    if name_idx == -1 or class_idx == -1 or roll_idx == -1:
        return {
            "success": False,
            "message": "Required columns missing! File must have at least: 'Full Name', 'Class', and 'Roll No'."
        }

    conn = get_db_connection()
    cursor = conn.cursor()

    added = 0
    updated = 0
    now_iso = datetime.now(PKT).isoformat()

    for row_num, row in enumerate(data_rows, start=2):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue

        name = str(row[name_idx] or "").strip()
        class_name = str(row[class_idx] or "").strip()
        roll_no = str(row[roll_idx] or "").strip()

        if not name or not class_name or not roll_no:
            errors.append(f"Row {row_num}: Missing Name, Class, or Roll No.")
            continue

        father_name = str(row[father_idx] or "").strip() if father_idx != -1 and len(row) > father_idx else ""
        parent_phone = str(row[phone_idx] or "").strip() if phone_idx != -1 and len(row) > phone_idx else ""
        gender = str(row[gender_idx] or "Male").strip() if gender_idx != -1 and len(row) > gender_idx else "Male"
        address = str(row[address_idx] or "").strip() if address_idx != -1 and len(row) > address_idx else ""

        # Always compute Option-B code: S{class}{roll:02d}
        student_code, clean_roll = compute_student_code_and_roll(class_name, roll_no)

        # Check if student with same student_code already exists
        existing = cursor.execute(
            "SELECT id, student_code FROM students WHERE student_code = ? OR (class_name = ? AND roll_no = ?)",
            (student_code, class_name, clean_roll)
        ).fetchone()

        if existing:
            # Update existing
            cursor.execute("""
                UPDATE students 
                SET full_name = ?, father_name = ?, parent_phone = ?, gender = ?, address = ?, is_active = 1
                WHERE id = ?
            """, (name, father_name, parent_phone, gender, address, existing["id"]))
            updated += 1
        else:
            # Insert new with Option-B code
            qr_payload = generate_qr_payload(student_code)
            cursor.execute("""
                INSERT INTO students (student_code, full_name, father_name, class_name, roll_no, parent_phone, gender, address, qr_payload, is_active, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?)
            """, (student_code, name, father_name, class_name, clean_roll, parent_phone, gender, address, qr_payload, now_iso))
            added += 1

    conn.commit()
    conn.close()

    return {
        "success": True,
        "added": added,
        "updated": updated,
        "errors": errors,
        "message": f"Successfully imported: {added} new students added, {updated} updated."
    }

def export_attendance_to_excel(
    school_name: str,
    start_date: str,
    end_date: str,
    class_filter: str = ""
) -> io.BytesIO:
    """Generates comprehensive Excel report with summary sheet and detailed records sheet"""
    wb = openpyxl.Workbook()
    
    # ------------------ SHEET 1: SUMMARY ------------------
    ws_summary = wb.active
    ws_summary.title = "Attendance Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    # Header styling
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="475569")
    section_font = Font(name="Calibri", size=12, bold=True, color="0F172A")
    th_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    th_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    accent_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    ws_summary.append([school_name])
    ws_summary.cell(row=1, column=1).font = title_font
    ws_summary.append([f"Attendance Report: {start_date} to {end_date} | Class: {class_filter or 'All Classes'}"])
    ws_summary.cell(row=2, column=1).font = subtitle_font
    ws_summary.append([]) # blank

    # Fetch aggregate data
    conn = get_db_connection()
    cursor = conn.cursor()

    query_counts = """
        SELECT 
            COUNT(DISTINCT s.id) as total_students,
            COUNT(CASE WHEN a.status = 'Present' THEN 1 END) as present_count,
            COUNT(CASE WHEN a.status = 'Late' THEN 1 END) as late_count,
            COUNT(CASE WHEN a.status = 'Absent' THEN 1 END) as absent_count,
            COUNT(CASE WHEN a.status = 'Leave' THEN 1 END) as leave_count
        FROM students s
        LEFT JOIN attendance a ON s.id = a.student_id AND a.date >= ? AND a.date <= ?
        WHERE s.is_active = 1
    """
    params_counts = [start_date, end_date]
    if class_filter:
        query_counts += " AND s.class_name = ?"
        params_counts.append(class_filter)

    totals = dict(cursor.execute(query_counts, params_counts).fetchone())

    # Summary table
    ws_summary.append(["METRIC", "COUNT"])
    for col_idx in [1, 2]:
        cell = ws_summary.cell(row=4, column=col_idx)
        cell.font = th_font
        cell.fill = th_fill

    ws_summary.append(["Total Active Students", totals.get("total_students", 0)])
    ws_summary.append(["Total Present Marks", totals.get("present_count", 0)])
    ws_summary.append(["Total Late Marks", totals.get("late_count", 0)])
    ws_summary.append(["Total Absent Marks", totals.get("absent_count", 0)])
    ws_summary.append(["Total Approved Leave", totals.get("leave_count", 0)])

    # Class-wise Breakdown
    ws_summary.append([])
    ws_summary.append(["CLASS-WISE BREAKDOWN"])
    ws_summary.cell(row=11, column=1).font = section_font

    class_headers = ["Class", "Active Students", "Present", "Late", "Absent", "Leave", "Attendance %"]
    ws_summary.append(class_headers)
    for col_idx in range(1, len(class_headers) + 1):
        cell = ws_summary.cell(row=12, column=col_idx)
        cell.font = th_font
        cell.fill = th_fill

    class_query = """
        SELECT 
            s.class_name,
            COUNT(DISTINCT s.id) as class_students,
            COUNT(CASE WHEN a.status = 'Present' THEN 1 END) as p,
            COUNT(CASE WHEN a.status = 'Late' THEN 1 END) as l,
            COUNT(CASE WHEN a.status = 'Absent' THEN 1 END) as ab,
            COUNT(CASE WHEN a.status = 'Leave' THEN 1 END) as lv
        FROM students s
        LEFT JOIN attendance a ON s.id = a.student_id AND a.date >= ? AND a.date <= ?
        WHERE s.is_active = 1
    """
    class_params = [start_date, end_date]
    if class_filter:
        class_query += " AND s.class_name = ?"
        class_params.append(class_filter)
    class_query += " GROUP BY s.class_name ORDER BY s.class_name ASC"

    class_rows = cursor.execute(class_query, class_params).fetchall()
    for row in class_rows:
        stu_cnt = row["class_students"]
        p = row["p"]
        l = row["l"]
        tot_marked = p + l + row["ab"] + row["lv"]
        rate = f"{round(((p + l) / tot_marked * 100), 1)}%" if tot_marked > 0 else "0%"
        ws_summary.append([row["class_name"], stu_cnt, p, l, row["ab"], row["lv"], rate])

    # ------------------ SHEET 2: DETAILED LOG ------------------
    ws_detail = wb.create_sheet(title="Detailed Attendance Log")
    ws_detail.views.sheetView[0].showGridLines = True

    detail_headers = [
        "Date", "Student Code", "Student Name", "Father Name", "Class", "Roll No", "Status", "Time", "Method", "Notes"
    ]
    ws_detail.append(detail_headers)
    for col_idx in range(1, len(detail_headers) + 1):
        cell = ws_detail.cell(row=1, column=col_idx)
        cell.font = th_font
        cell.fill = th_fill

    log_query = """
        SELECT 
            a.date,
            s.student_code,
            s.full_name,
            s.father_name,
            s.class_name,
            s.roll_no,
            a.status,
            COALESCE(a.scan_time, '-') as scan_time,
            a.scan_method,
            COALESCE(a.notes, '') as notes
        FROM attendance a
        JOIN students s ON a.student_id = s.id
        WHERE a.date >= ? AND a.date <= ?
    """
    log_params = [start_date, end_date]
    if class_filter:
        log_query += " AND s.class_name = ?"
        log_params.append(class_filter)
    log_query += " ORDER BY a.date DESC, s.class_name ASC, CAST(s.roll_no AS INTEGER) ASC"

    records = cursor.execute(log_query, log_params).fetchall()
    for r in records:
        ws_detail.append([
            r["date"], r["student_code"], r["full_name"], r["father_name"],
            r["class_name"], r["roll_no"], r["status"], r["scan_time"],
            r["scan_method"], r["notes"]
        ])

    conn.close()

    # Column autosizing for both sheets
    for ws in [ws_summary, ws_detail]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
